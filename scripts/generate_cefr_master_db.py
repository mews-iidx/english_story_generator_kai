import os
import re
import json
import string
import urllib.request
import requests
from bs4 import BeautifulSoup
import openpyxl

PROJECT_ROOT = "/home/mews/english_storie_generator_kai"
EGP_FILE = os.path.join(PROJECT_ROOT, "egpo.xlsx")
VOCAB_OUTPUT = os.path.join(PROJECT_ROOT, "src/data/cefrVocabMaster.ts")
PATTERNS_OUTPUT = os.path.join(PROJECT_ROOT, "src/data/cefrPatternsMaster.ts")
TYPES_OUTPUT = os.path.join(PROJECT_ROOT, "src/types/mastery.ts")
SCRIPTS_DIR = os.path.join(PROJECT_ROOT, "scripts")
os.makedirs(SCRIPTS_DIR, exist_ok=True)

# 1. Download Cambridge EGP if missing
if not os.path.exists(EGP_FILE):
    print("Downloading Cambridge EGP...")
    egp_url = "https://raw.githubusercontent.com/ninja33/EGP/master/asset/egpo.xlsx"
    urllib.request.urlretrieve(egp_url, EGP_FILE)
    print("Cambridge EGP downloaded.")

# 2. Download / Load EJDict for Japanese translations
print("Loading EJDict for Japanese translations...")
ejdict = {}
for ch in string.ascii_lowercase:
    url = f"https://raw.githubusercontent.com/kujirahand/EJDict/master/src/{ch}.txt"
    r = requests.get(url)
    if r.status_code == 200:
        for line in r.text.splitlines():
            if '\t' in line:
                k, v = line.split('\t', 1)
                k_clean = k.strip().lower()
                if k_clean not in ejdict:
                    ejdict[k_clean] = v.strip()

print(f"Loaded {len(ejdict)} EJDict entries.")

def clean_ejdict_meaning(raw_meaning):
    if not raw_meaning:
        return ""
    cleaned = re.sub(r'〈[^〉]+〉', '', raw_meaning)
    cleaned = re.sub(r'《[^》]+》', '', cleaned)
    cleaned = re.sub(r'\[[^\]]+\]', '', cleaned)
    cleaned = re.sub(r'\([^\)]+\)', '', cleaned)
    parts = [p.strip() for p in re.split(r'[/,;、]', cleaned) if p.strip()]
    if parts:
        res = ", ".join(parts[:2])
        return res
    return raw_meaning[:40].strip()

# 3. Scrape Oxford 3000/5000
print("Scraping Oxford 3000/5000 wordlist...")
headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
ox_url = "https://www.oxfordlearnersdictionaries.com/wordlists/oxford3000-5000"
res = requests.get(ox_url, headers=headers)
soup = BeautifulSoup(res.text, 'html.parser')
word_items = soup.find_all('li', {'data-hw': True})
print(f"Found {len(word_items)} raw items on Oxford page.")

pos_ja_map = {
    'noun': '名詞',
    'verb': '動詞',
    'adjective': '形容詞',
    'adverb': '副詞',
    'preposition': '前置詞',
    'conjunction': '接続詞',
    'pronoun': '代名詞',
    'determiner': '限定詞',
    'modal verb': '助動詞',
    'auxiliary verb': '助動詞',
    'number': '数詞',
    'exclamation': '感動詞',
    'indefinite article': '不定冠詞',
    'definite article': '定冠詞',
    'linking verb': '連結動詞',
}

fallback_meanings = {
    'according to': '〜によれば、〜に従って',
    'accountability': '説明責任、責務',
    'affordable': '手頃な、購入しやすい',
    'afterwards': 'その後、後で',
    'aids': 'エイズ (後天性免疫不全症候群)',
    'align': '整列させる、一致させる',
    'alignment': '整列、連携',
    'all right': '大丈夫な、了解して',
    'annoyed': 'いらだった、困惑した',
    'any more': 'これ以上、もはや',
    'app': 'アプリ、アプリケーション',
    'archive': '保存記録、保管する',
    'arguably': 'ほぼ間違いなく、議論の余地はあるが',
    'artwork': '芸術作品、挿絵',
    'associated': '関連した、提携した',
    'backup': '予備、バックアップ',
    'based': '〜に基づいた',
    'behaviour': '行動、振る舞い',
    'best-selling': 'ベストセラーの',
    'biodiversity': '生物多様性',
    'bizarre': '奇妙な、風変わりな',
    'bleeding': '出血',
    'boredom': '退屈',
    'broadband': 'ブロードバンド高速通信',
    'browser': 'ブラウザ、閲覧ソフト',
    'by means of': '〜によって',
    'carer': '介護者、世話人',
    'celebrated': '有名な、著名な',
    'centre': '中心、センター',
    'colour': '色',
    'favour': '好意、恩恵',
    'flavour': '風味、味',
    'honour': '名誉、栄誉',
    'humour': 'ユーモア',
    'labour': '労働、勤労',
    'neighbour': '隣人、近所の人',
    'rumour': '噂、風評',
    'the': 'その (定冠詞)',
    'a': 'ある、1つの (不定冠詞)',
    'an': 'ある、1つの (不定冠詞)',
    'per cent': 'パーセント',
    'in order to': '〜するために',
    'as well as': '〜と同様に',
    'such as': '〜のような',
    'so that': '〜できるように',
    'at least': '少なくとも',
    'of course': 'もちろん',
    'each other': 'お互い',
    'as soon as': '〜するとすぐに',
    'as long as': '〜である限り',
}

vocab_records = []
seen_words = set()

for item in word_items:
    hw = item.get('data-hw', '').strip()
    cefr = (item.get('data-ox5000') or item.get('data-ox3000') or '').upper()
    if cefr not in ['A1', 'A2', 'B1', 'B2']:
        continue
    
    pos_span = item.find('span', class_='pos')
    pos_raw = pos_span.text.strip().lower() if pos_span else 'noun'
    pos_ja = pos_ja_map.get(pos_raw, pos_raw)

    dedup_key = f"{hw.lower()}_{pos_raw}"
    if dedup_key in seen_words:
        continue
    seen_words.add(dedup_key)

    hw_lower = hw.lower()
    meaning = ""
    if hw_lower in fallback_meanings:
        meaning = fallback_meanings[hw_lower]
    elif hw_lower in ejdict:
        meaning = clean_ejdict_meaning(ejdict[hw_lower])
    else:
        alt1 = re.sub(r'our$', 'or', hw_lower)
        alt2 = re.sub(r'ise$', 'ize', hw_lower)
        alt3 = re.sub(r're$', 'er', hw_lower)
        if alt1 in ejdict:
            meaning = clean_ejdict_meaning(ejdict[alt1])
        elif alt2 in ejdict:
            meaning = clean_ejdict_meaning(ejdict[alt2])
        elif alt3 in ejdict:
            meaning = clean_ejdict_meaning(ejdict[alt3])
        elif ' ' in hw_lower:
            words = hw_lower.split()
            meanings = [clean_ejdict_meaning(ejdict.get(w, '')) for w in words if ejdict.get(w)]
            meaning = " ".join([m for m in meanings if m])
        if not meaning:
            meaning = hw

    vocab_records.append({
        'phrase': hw,
        'meaning': meaning,
        'cefr': cefr,
        'partOfSpeech': pos_ja,
    })

vocab_by_level = {'A1': [], 'A2': [], 'B1': [], 'B2': []}
for r in vocab_records:
    vocab_by_level[r['cefr']].append(r)

final_vocabs_by_level = {'A1': [], 'A2': [], 'B1': [], 'B2': []}
for lvl in ['A1', 'A2', 'B1', 'B2']:
    for idx, r in enumerate(vocab_by_level[lvl], 1):
        final_vocabs_by_level[lvl].append({
            'id': f"vocab_{lvl.lower()}_{idx:04d}",
            'phrase': r['phrase'],
            'meaning': r['meaning'],
            'cefr': lvl,
            'partOfSpeech': r['partOfSpeech'],
        })

print(f"Final Vocab Count: A1={len(final_vocabs_by_level['A1'])}, A2={len(final_vocabs_by_level['A2'])}, B1={len(final_vocabs_by_level['B1'])}, B2={len(final_vocabs_by_level['B2'])}")

# 4. Parse Cambridge English Grammar Profile
print("Parsing Cambridge English Grammar Profile...")
wb = openpyxl.load_workbook(EGP_FILE)
ws = wb.active

category_map = {
    'ADJECTIVES': ('adjectives_adverbs', '形容詞・限定表現'),
    'ADVERBS': ('adjectives_adverbs', '副詞・修飾表現'),
    'CLAUSES': ('relative_clauses', '節・複文・関係節'),
    'CONJUNCTIONS': ('coordination', '等位・従属接続詞'),
    'DETERMINERS': ('determiners', '冠詞・限定詞'),
    'DISCOURSE MARKERS': ('coordination', '談話標識・論理マーカー'),
    'FOCUS': ('inversion_emphasis', '強調・倒置構文'),
    'FUTURE': ('tenses_aspects', '未来表現・時制'),
    'MODALITY': ('modals', '助動詞・モダリティ'),
    'NEGATION': ('questions_negatives', '否定・部分否定'),
    'NOUNS': ('nouns_pronouns', '名詞・代名詞'),
    'PASSIVES': ('causative_passive', '受動態・使役構文'),
    'PAST': ('tenses_aspects', '過去時制・完了形'),
    'PREPOSITIONS': ('prepositions', '前置詞・前置詞句'),
    'PRESENT': ('tenses_aspects', '現在時制・進行形'),
    'PRONOUNS': ('nouns_pronouns', '代名詞・照応表現'),
    'QUESTIONS': ('questions_negatives', '疑問文・付加疑問'),
    'REPORTED SPEECH': ('reported_speech', '話法・間接話法'),
    'VERBS': ('infinitives_gerunds', '動詞・不定詞・動名詞'),
}

def clean_example_text(raw_text):
    if not raw_text:
        return []
    blocks = [b.strip() for b in re.split(r'\n+', raw_text) if b.strip()]
    cleaned_sentences = []
    for b in blocks:
        cleaned = re.sub(r'\([^\)]*(?:Pass|THRESHOLD|BREAKTHROUGH|WAYSTAGE|VANTAGE|MASTERY|Polish|Spanish|German|Japan|China|French|Italian|Russian|Greek|Turkish|200\d|199\d)[^\)]*\)', '', b)
        cleaned = re.sub(r'\[[^\]]*\]', '', cleaned)
        cleaned = re.sub(r'\s+', ' ', cleaned).strip()
        cleaned = re.sub(r'\s+([,\.\?!;:])', r'\1', cleaned)
        if len(cleaned) > 5 and not cleaned.startswith('►'):
            if not cleaned[-1] in '.?!':
                cleaned += '.'
            cleaned_sentences.append(cleaned)
    return cleaned_sentences

def extract_target_tokens(sentence, guideword, can_do):
    tokens = []
    quoted = re.findall(r"'([^']+)'", guideword + " " + can_do)
    for q in quoted:
        q_clean = q.strip().lower()
        if len(q_clean) > 0 and q_clean in sentence.lower():
            tokens.append(q_clean)
    
    if not tokens:
        words = re.findall(r'\b[a-zA-Z]+\b', sentence)
        for w in words:
            if w.lower() in ['can', 'could', 'would', 'should', 'might', 'must', 'have', 'had', 'been', 'is', 'are', 'was', 'were', 'if', 'because', 'although', 'which', 'who', 'that', 'than', 'as', 'to', 'for', 'with', 'by']:
                tokens.append(w.lower())
                break
    if not tokens and len(words) >= 2:
        tokens = [words[0].lower(), words[1].lower()]
    return list(dict.fromkeys(tokens))[:2]

pattern_records = []
for r in range(2, ws.max_row + 1):
    super_cat = ws.cell(r, 2).value
    sub_cat = ws.cell(r, 3).value
    lvl = ws.cell(r, 4).value
    guideword = ws.cell(r, 6).value or ''
    can_do = ws.cell(r, 7).value or ''
    raw_examples = ws.cell(r, 8).value or ''

    if lvl not in ['A1', 'A2', 'B1', 'B2']:
        continue

    cat_key, cat_label = category_map.get(super_cat, ('basic_syntax', '基本構文'))
    sents = clean_example_text(str(raw_examples))

    while len(sents) < 3:
        if len(sents) == 2:
            sents.append(sents[0])
        elif len(sents) == 1:
            sents.append(sents[0])
            sents.append(sents[0])
        else:
            sents = [
                "This is an example sentence for this grammar structure.",
                "Students practice this pattern in everyday English conversations.",
                "You can express clear ideas using this grammatical form."
            ]

    name_clean = guideword.strip() if guideword else f"{super_cat}: {sub_cat}"
    name_display = re.sub(r'^(?:FORM|USE|FORM/USE|MEANING):\s*', '', name_clean)
    if not name_display:
        name_display = f"{super_cat} - {sub_cat}"

    variations = []
    for s in sents[:3]:
        target_tokens = extract_target_tokens(s, guideword, can_do)
        variations.append({
            'sentence': s,
            'translation': f"【例文】{s}",
            'targetTokens': target_tokens
        })

    pattern_records.append({
        'cefr': lvl,
        'category': cat_key,
        'categoryLabel': cat_label,
        'name': name_display[:80],
        'meaning': can_do.strip(),
        'focus': f"{super_cat} / {sub_cat}: {guideword}".strip(),
        'variations': variations
    })

patterns_by_level = {'A1': [], 'A2': [], 'B1': [], 'B2': []}
for p in pattern_records:
    patterns_by_level[p['cefr']].append(p)

final_patterns_by_level = {'A1': [], 'A2': [], 'B1': [], 'B2': []}
for lvl in ['A1', 'A2', 'B1', 'B2']:
    for idx, p in enumerate(patterns_by_level[lvl], 1):
        final_patterns_by_level[lvl].append({
            'id': f"pat_{lvl.lower()}_{idx:04d}",
            'cefr': lvl,
            'category': p['category'],
            'categoryLabel': p['categoryLabel'],
            'name': p['name'],
            'meaning': p['meaning'],
            'focus': p['focus'],
            'variations': p['variations']
        })

print(f"Final Patterns Count: A1={len(final_patterns_by_level['A1'])}, A2={len(final_patterns_by_level['A2'])}, B1={len(final_patterns_by_level['B1'])}, B2={len(final_patterns_by_level['B2'])}")

# 5. Output Types Update
print("Updating types/mastery.ts...")
with open(TYPES_OUTPUT, 'r', encoding='utf-8') as f:
    mastery_ts_code = f.read()

pattern_cat_definition = """export type PatternCategory =
  | 'basic_syntax'
  | 'infinitives_gerunds'
  | 'tenses_aspects'
  | 'modals'
  | 'questions_negatives'
  | 'coordination'
  | 'correlative'
  | 'relative_clauses'
  | 'participles'
  | 'comparatives'
  | 'conditionals'
  | 'inversion_emphasis'
  | 'causative_passive'
  | 'adjectives_adverbs'
  | 'nouns_pronouns'
  | 'prepositions'
  | 'determiners'
  | 'reported_speech';"""

mastery_ts_code = re.sub(r'export type PatternCategory =[\s\S]*?;\n', pattern_cat_definition + "\n", mastery_ts_code, count=1)
with open(TYPES_OUTPUT, 'w', encoding='utf-8') as f:
    f.write(mastery_ts_code)

# 6. Write src/data/cefrVocabMaster.ts
print("Writing src/data/cefrVocabMaster.ts...")
vocab_ts_content = f"""import {{ VocabMasterItem }} from '../types/mastery';
import {{ CefrLevel }} from '../types/settings';

export const A1_VOCABS: VocabMasterItem[] = {json.dumps(final_vocabs_by_level['A1'], ensure_ascii=False, indent=2)};
export const A2_VOCABS: VocabMasterItem[] = {json.dumps(final_vocabs_by_level['A2'], ensure_ascii=False, indent=2)};
export const B1_VOCABS: VocabMasterItem[] = {json.dumps(final_vocabs_by_level['B1'], ensure_ascii=False, indent=2)};
export const B2_VOCABS: VocabMasterItem[] = {json.dumps(final_vocabs_by_level['B2'], ensure_ascii=False, indent=2)};

export const VOCABS_BY_LEVEL: Record<CefrLevel, VocabMasterItem[]> = {{
  A1: A1_VOCABS,
  A2: A2_VOCABS,
  B1: B1_VOCABS,
  B2: B2_VOCABS,
  C1: [],
}};

export const CEFR_VOCAB_MASTER: VocabMasterItem[] = [
  ...A1_VOCABS,
  ...A2_VOCABS,
  ...B1_VOCABS,
  ...B2_VOCABS,
];

export function getVocabMasterByLevel(level: CefrLevel): VocabMasterItem[] {{
  return VOCABS_BY_LEVEL[level] || [];
}}

export function getVocabById(id: string): VocabMasterItem | undefined {{
  return CEFR_VOCAB_MASTER.find(v => v.id === id);
}}

export function getVocabByPhrase(phrase: string): VocabMasterItem | undefined {{
  const lower = phrase.trim().toLowerCase();
  return CEFR_VOCAB_MASTER.find(v => v.phrase.toLowerCase() === lower);
}}

export function getRandomVocabsByLevel(level: CefrLevel, count: number = 5): VocabMasterItem[] {{
  const pool = getVocabMasterByLevel(level);
  if (pool.length <= count) return pool;
  const shuffled = [...pool].sort(() => 0.5 - Math.random());
  return shuffled.slice(0, count);
}}
"""

with open(VOCAB_OUTPUT, 'w', encoding='utf-8') as f:
    f.write(vocab_ts_content)

# 7. Write src/data/cefrPatternsMaster.ts
print("Writing src/data/cefrPatternsMaster.ts...")
patterns_ts_content = f"""import {{ PatternMasterItem }} from '../types/mastery';
import {{ CefrLevel }} from '../types/settings';

export const A1_PATTERNS: PatternMasterItem[] = {json.dumps(final_patterns_by_level['A1'], ensure_ascii=False, indent=2)};
export const A2_PATTERNS: PatternMasterItem[] = {json.dumps(final_patterns_by_level['A2'], ensure_ascii=False, indent=2)};
export const B1_PATTERNS: PatternMasterItem[] = {json.dumps(final_patterns_by_level['B1'], ensure_ascii=False, indent=2)};
export const B2_PATTERNS: PatternMasterItem[] = {json.dumps(final_patterns_by_level['B2'], ensure_ascii=False, indent=2)};

export const PATTERNS_BY_LEVEL: Record<CefrLevel, PatternMasterItem[]> = {{
  A1: A1_PATTERNS,
  A2: A2_PATTERNS,
  B1: B1_PATTERNS,
  B2: B2_PATTERNS,
  C1: [],
}};

export const CEFR_PATTERNS_MASTER: PatternMasterItem[] = [
  ...A1_PATTERNS,
  ...A2_PATTERNS,
  ...B1_PATTERNS,
  ...B2_PATTERNS,
];

export function getPatternsByLevel(level: CefrLevel): PatternMasterItem[] {{
  return PATTERNS_BY_LEVEL[level] || [];
}}

export function getPatternById(id: string): PatternMasterItem | undefined {{
  return CEFR_PATTERNS_MASTER.find(p => p.id === id);
}}

export function getRandomPatternsByLevel(level: CefrLevel, count: number = 3): PatternMasterItem[] {{
  const pool = getPatternsByLevel(level);
  if (pool.length <= count) return pool;
  const shuffled = [...pool].sort(() => 0.5 - Math.random());
  return shuffled.slice(0, count);
}}
"""

with open(PATTERNS_OUTPUT, 'w', encoding='utf-8') as f:
    f.write(patterns_ts_content)

# Also save a copy of this script in project root scripts/
script_in_repo = os.path.join(SCRIPTS_DIR, "generate_cefr_master_db.py")
with open(script_in_repo, 'w', encoding='utf-8') as f:
    with open(__file__, 'r', encoding='utf-8') as src_f:
        f.write(src_f.read())

print("Successfully generated and updated all CEFR Master files!")
